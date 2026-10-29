import io

import os

import re

from datetime import datetime

import sqlite3

import json

import joblib

import pandas as pd

import pytesseract

import pymupdf

from docx import Document

from flask import Flask, request, jsonify

from flask_cors import CORS

from PIL import Image, ImageOps

from werkzeug.utils import secure_filename

# --------------------------------------------------

# CONFIGURATION

# --------------------------------------------------

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

MODEL_PATH = os.path.join(BASE_DIR, "shortlisting_model.pkl")

TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

if os.path.isfile(TESSERACT_PATH):

    pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH

if not os.path.isfile(MODEL_PATH):

    raise FileNotFoundError(f"Trained model not found: {MODEL_PATH}")

# Load the existing trained model. No retraining.

model = joblib.load(MODEL_PATH)

FEATURES = list(model.feature_names_in_)

EXPECTED_FEATURES = [

    "age", "education_level", "university_tier", "cgpa", "internships",

    "projects", "programming_languages", "certifications", "experience_years",

    "hackathons", "research_papers", "skills_score", "soft_skills_score",

    "resume_length_words", "company_type",

]

if FEATURES != EXPECTED_FEATURES:

    raise RuntimeError(

        "Model feature schema mismatch. Check the trained model's feature order."

    )

preprocessor = model.named_steps["preprocessor"]

NUMERIC_FIELDS = [

    "age", "cgpa", "internships", "projects", "programming_languages",

    "certifications", "experience_years", "hackathons", "research_papers",

    "skills_score", "soft_skills_score", "resume_length_words",

]

CATEGORICAL_FIELDS = ["education_level", "university_tier", "company_type"]

actual_numeric_fields = next(

    columns for name, transformer, columns in preprocessor.transformers_

    if name == "num"

)

actual_categorical_fields = next(

    columns for name, transformer, columns in preprocessor.transformers_

    if name == "cat"

)

if set(NUMERIC_FIELDS) != set(actual_numeric_fields):

    raise RuntimeError("Numeric fields do not match the trained pipeline.")

if set(CATEGORICAL_FIELDS) != set(actual_categorical_fields):

    raise RuntimeError("Categorical fields do not match the trained pipeline.")

app = Flask(__name__)

CORS(app)

# --------------------------------------------------

# SCREENING HISTORY DATABASE

# --------------------------------------------------

DATABASE_PATH = os.path.join(BASE_DIR, "screening_history.db")

def get_db_connection():

    connection = sqlite3.connect(DATABASE_PATH)

    connection.row_factory = sqlite3.Row

    return connection

def init_db():

    connection = get_db_connection()

    try:

        connection.execute("""

            CREATE TABLE IF NOT EXISTS screening_history (

                id INTEGER PRIMARY KEY AUTOINCREMENT,

                candidate_data TEXT NOT NULL,

                prediction TEXT NOT NULL,

                result TEXT NOT NULL,

                shortlisted_probability REAL,

                created_at TEXT NOT NULL,

                candidate_name TEXT,

                resume_filename TEXT

            )

        """)

        # Safely migrate an existing database without deleting records.

        columns = {

            row["name"]

            for row in connection.execute(

                "PRAGMA table_info(screening_history)"

            ).fetchall()

        }

        if "candidate_name" not in columns:

            connection.execute(

                "ALTER TABLE screening_history ADD COLUMN candidate_name TEXT"

            )

        if "resume_filename" not in columns:

            connection.execute(

                "ALTER TABLE screening_history ADD COLUMN resume_filename TEXT"

            )

        connection.commit()

    except sqlite3.Error:

        connection.rollback()

        app.logger.exception("Database initialization failed")

        raise

    finally:

        connection.close()

init_db()

app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024

ALLOWED_EXTENSIONS = {"pdf", "docx", "jpg", "jpeg", "png", "webp"}

RELEVANT_SKILLS = {

    "java", "python", "javascript", "c++", "c", "html", "css", "react",

    "node.js", "express", "sql", "mysql", "mongodb", "rest api", "git",

    "github", "data structures", "algorithms", "oops", "machine learning",

    "scikit-learn", "flask", "api integration", "bootstrap", "tailwind",

    "communication", "problem solving",

}

# --------------------------------------------------

# GENERAL TEXT / SECTION HELPERS

# --------------------------------------------------

def normalize_text(text):

    text = (text or "").replace("\r\n", "\n")

    text = text.replace("\r", "\n").replace("\u00a0", " ")

    text = re.sub(r"[ \t]+", " ", text)

    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()

def normalize_heading(line):

    """Normalize a heading while tolerating PDF bullets, numbering and punctuation."""

    cleaned = line.strip().upper()

    cleaned = re.sub(r"^[\s#>*•●▪\d.)-]+", "", cleaned)

    cleaned = re.sub(r"[:|]+$", "", cleaned)

    cleaned = re.sub(r"[^A-Z& /]", "", cleaned)

    cleaned = re.sub(r"\s+", " ", cleaned)

    return cleaned.strip()

def section_lines(text, start_headings, stop_headings):

    """Return lines between a recognized section heading and the next section heading."""

    collected = []

    in_section = False

    for line in text.splitlines():

        heading = normalize_heading(line)

        if heading in start_headings:

            in_section = True

            continue

        if in_section and heading in stop_headings:

            break

        if in_section:

            collected.append(line)

    return collected

def calculate_skills_score(text):

    """Keyword coverage only; this is not a measure of proficiency."""

    normalized = text.lower()

    found = set()

    for skill in RELEVANT_SKILLS:

        pattern = (

            r"(?<![a-z0-9+#.])" + re.escape(skill.lower())

            + r"(?![a-z0-9+#.])"

        )

        if re.search(pattern, normalized):

            found.add(skill)

    if not RELEVANT_SKILLS:

        return 0.0

    return round(len(found) / len(RELEVANT_SKILLS) * 100, 2)

# --------------------------------------------------

# EXPERIENCE / INTERNSHIP / PROJECT PARSING

# --------------------------------------------------

EXPERIENCE_HEADINGS = {

    "EXPERIENCE", "WORK EXPERIENCE", "PROFESSIONAL EXPERIENCE",

    "INTERNSHIP", "INTERNSHIPS", "WORK HISTORY", "EMPLOYMENT HISTORY",

    "EXPERIENCE & INTERNSHIPS", "INTERNSHIP EXPERIENCE",

}

PROJECT_HEADINGS = {

    "PROJECTS", "PERSONAL PROJECTS", "ACADEMIC PROJECTS",

    "KEY PROJECTS", "PROJECT EXPERIENCE", "TECHNICAL PROJECTS",

}

CERTIFICATE_HEADINGS = {

    "CERTIFICATIONS", "CERTIFICATES", "LICENSES AND CERTIFICATIONS",

}

COMMON_STOP_HEADINGS = {

    "EDUCATION", "ACADEMIC QUALIFICATIONS", "QUALIFICATIONS",

    "PROJECTS", "PERSONAL PROJECTS", "ACADEMIC PROJECTS",

    "KEY PROJECTS", "PROJECT EXPERIENCE", "TECHNICAL PROJECTS",

    "TECHNICAL SKILLS", "SKILLS", "CERTIFICATIONS", "CERTIFICATES",

    "LICENSES AND CERTIFICATIONS", "ACHIEVEMENTS",

    "LEADERSHIP & ACHIEVEMENTS", "PUBLICATIONS", "INTERESTS",

    "LANGUAGES", "SUMMARY", "PROFESSIONAL SUMMARY", "PROFILE",

    "CONTACT", "PERSONAL DETAILS", "HACKATHON", "HACKATHONS",

    "HACKATHON EXPERIENCE", "HACKATHONS AND COMPETITIONS",

    "COMPETITIONS AND HACKATHONS", "RESEARCH", "RESEARCH PAPERS",

    "RESEARCH AND PUBLICATIONS", "PUBLICATIONS AND RESEARCH", "PAPERS",

    "SELECTED PUBLICATIONS",

}

def extract_experience_years(text):

    """

    Estimate experience from dates in experience/work/internship sections only.

    Month ranges are inclusive; overlapping months count once.

    Standalone month-year entries count as one month and should be reviewed.

    """

    lines = section_lines(

        text,

        EXPERIENCE_HEADINGS,

        COMMON_STOP_HEADINGS - EXPERIENCE_HEADINGS,

    )

    # Some resumes put Experience and Projects next to each other without

    # a standalone heading; only text inside a recognized experience section is used.

    section_text = "\n".join(lines)

    if not section_text.strip():

        return None

    month_pattern = (

        r"Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|"

        r"Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|"

        r"Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?"

    )

    month_numbers = {

        "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,

        "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,

    }

    def month_index(month, year):

        return int(year) * 12 + month_numbers[month[:3].lower()] - 1

    now = datetime.now()

    current_index = now.year * 12 + now.month - 1

    year_pattern = r"(?:19|20)\d{2}"

    range_pattern = re.compile(

        rf"\b({month_pattern})\s+({year_pattern})"

        rf"\s*(?:-|–|—|\bto\b)\s*"

        rf"({month_pattern}\s+{year_pattern}|Present|Current|Till\s+Date)\b",

        re.IGNORECASE,

    )

    standalone_pattern = re.compile(

        rf"\b({month_pattern})\s+({year_pattern})\b", re.IGNORECASE

    )

    intervals = []

    covered_spans = []

    for match in range_pattern.finditer(section_text):

        start_month, start_year = match.group(1), match.group(2)

        end_text = match.group(3).strip()

        start = month_index(start_month, start_year)

        if end_text.lower() in {"present", "current", "till date"}:

            end = current_index

        else:

            end_match = re.fullmatch(

                rf"({month_pattern})\s+({year_pattern})",

                end_text,

                re.IGNORECASE,

            )

            if not end_match:

                continue

            end = month_index(end_match.group(1), end_match.group(2))

        # Ignore future start dates and invalid/reversed ranges.

        if start > current_index or end < start:

            covered_spans.append(match.span())

            continue

        # Use inclusive calendar months and cap ongoing roles at current month.

        intervals.append((start, min(end, current_index) + 1))

        covered_spans.append(match.span())

    # Handle a single date such as "Jul 2026" for a short internship.

    # Such entries count as one month and must be manually verified.

    for match in standalone_pattern.finditer(section_text):

        if any(

            span_start <= match.start() and match.end() <= span_end

            for span_start, span_end in covered_spans

        ):

            continue

        index = month_index(match.group(1), match.group(2))

        if index <= current_index:

            intervals.append((index, index + 1))

    if not intervals:

        return None

    # Merge overlapping/adjacent intervals so simultaneous roles aren't double-counted.

    intervals.sort()

    merged = []

    for start, end in intervals:

        if not merged or start > merged[-1][1]:

            merged.append([start, end])

        else:

            merged[-1][1] = max(merged[-1][1], end)

    total_months = sum(end - start for start, end in merged)

    return round(total_months / 12, 2)

HACKATHON_HEADINGS = {

    "HACKATHONS", "HACKATHON", "HACKATHON EXPERIENCE",

    "HACKATHONS AND COMPETITIONS", "COMPETITIONS AND HACKATHONS",

}

RESEARCH_HEADINGS = {

    "RESEARCH", "RESEARCH PAPERS", "PUBLICATIONS", "RESEARCH AND PUBLICATIONS",

    "PUBLICATIONS AND RESEARCH", "PAPERS", "SELECTED PUBLICATIONS",

}

TECHNICAL_SKILL_HEADINGS = {

    "TECHNICAL SKILLS", "TECHNICAL SKILL", "PROGRAMMING LANGUAGES",

    "TECHNOLOGIES", "TECHNOLOGY STACK", "SKILLS", "TECH STACK",

}

ENTRY_STOP_HEADINGS = {

    "EDUCATION", "ACADEMIC QUALIFICATIONS", "QUALIFICATIONS",

    "PROJECTS", "PERSONAL PROJECTS", "ACADEMIC PROJECTS", "KEY PROJECTS",

    "PROJECT EXPERIENCE", "TECHNICAL PROJECTS",

    "TECHNICAL SKILLS", "TECHNICAL SKILL", "PROGRAMMING LANGUAGES",

    "SKILLS", "CERTIFICATIONS", "CERTIFICATES",

    "LICENSES AND CERTIFICATIONS", "ACHIEVEMENTS",

    "LEADERSHIP & ACHIEVEMENTS", "PUBLICATIONS", "RESEARCH",

    "RESEARCH PAPERS", "HACKATHONS", "HACKATHON",

    "INTERESTS", "LANGUAGES", "SUMMARY", "PROFESSIONAL SUMMARY",

    "PROFILE", "CONTACT", "PERSONAL DETAILS",

}

BULLET_PREFIX = re.compile(r"^\s*(?:[•●▪◦*+-]+|\d+[.)])\s*")

PROSE_START = re.compile(

    r"^(?:built|developed|created|implemented|designed|worked on|"

    r"responsible for|contributed to|my role|description|"

    r"technologies used|technology stack|tech stack|tools used|"

    r"features include|key responsibilities|achievements?|"

    r"hackathon project|participated in|organized by|published in|"

    r"presented at|team size|secured rank|won|selected for)\b",

    re.I,

)

DATE_ONLY = re.compile(

    r"\b(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|"

    r"jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|"

    r"oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)\s+\d{4}\b",

    re.I,

)

def _clean_entry_line(raw_line):

    line = BULLET_PREFIX.sub("", raw_line or "").strip()

    # Remove common trailing metadata while preserving the title itself.

    line = re.sub(r"\s*[|•]\s*(?:tech(?:nologies)?|tools|stack)\s*:.+$", "", line, flags=re.I)

    return line.strip(" \t|:-–—")

def _is_likely_title(line, max_words=10):

    if not line:

        return False

    if len(line.split()) > max_words:

        return False

    if PROSE_START.search(line):

        return False

    if re.search(r"https?://|wwww\.", line, re.I):

        return False

    if line.endswith((".", "!", "?")):

        return False

    # Ignore obvious standalone date / duration / stack labels.

    if DATE_ONLY.fullmatch(line):

        return False

    if re.match(r"^(?:tech(?:nologies)?|tools|stack|duration|role|company|"

                r"responsibilities|description)\s*:", line, re.I):

        return False

    return bool(re.search(r"[A-Za-z]", line))

def _distinct_count(items):

    seen = set()

    for item in items:

        key = re.sub(r"[^a-z0-9]+", "", item.casefold())

        if key:

            seen.add(key)

    return len(seen)

def _is_project_technology_line(line):

    """Reject a technology-stack line that can look like a short project title."""

    cleaned = (line or "").strip().lower()

    if not cleaned:

        return False

    # Explicit stack labels are metadata, not project names.

    if re.match(r"^(?:tech(?:nologies)?(?: used)?|tools(?: used)?|stack)\s*[:|]", cleaned):

        return True

    # Common comma-separated technology lists, e.g. "React.js, Python, scikit-learn".

    known_tech = (

        "react", "react.js", "python", "java", "javascript", "html", "css",

        "scikit-learn", "node.js", "flask", "django", "fastapi", "mysql",

        "mongodb", "sql", "bootstrap", "tailwind", "tmdb api", "rest api",

        "c++", "c#", "typescript", "express", "pandas", "numpy",

    )

    parts = [part.strip(" .") for part in re.split(r"[,|]", cleaned) if part.strip()]

    if len(parts) >= 2:

        tech_hits = sum(

            any(part == tech or part.startswith(tech + " ") for tech in known_tech)

            for part in parts

        )

        if tech_hits >= 2 and tech_hits / len(parts) >= 0.6:

            return True

    # A short line consisting only of known technologies is also metadata.

    if len(parts) == 1 and any(cleaned == tech for tech in known_tech):

        return True

    return False

def extract_project_count(text):

    """Count project titles, ignoring stack lines and bullet descriptions."""

    lines = section_lines(

        text, PROJECT_HEADINGS, COMMON_STOP_HEADINGS - PROJECT_HEADINGS

    )

    if not lines:

        return None

    titles = []

    for raw_line in lines:

        original = (raw_line or "").strip()

        if not original:

            continue

        line = _clean_entry_line(original)

        if not line or _is_project_technology_line(line):

            continue

        # Description bullets should not be counted as separate projects.

        # Keep likely title bullets only when they are short and not action-led.

        is_bullet = bool(re.match(r"^[•●▪◦*+-]\s*", original))

        if is_bullet and PROSE_START.search(line):

            continue

        if re.search(r"\b(?:built|developed|created|implemented|designed|used|added|compared|classified|integrated|deployed|automated|enabled|allows|provides|supports)\b", line, re.I) and (line.endswith((".", "!", "?")) or len(line.split()) > 6):

            continue

        if not _is_likely_title(line, max_words=10):

            continue

        titles.append(line)

    count = _distinct_count(titles)

    return min(count, 50) if count else None

def extract_internship_count(text):

    """

    Count distinct internship role entries in experience sections.

    Count only short, non-bullet role/organization lines that identify an

    internship; repeated mentions in descriptive bullets are ignored.

    """

    lines = section_lines(

        text, EXPERIENCE_HEADINGS, COMMON_STOP_HEADINGS - EXPERIENCE_HEADINGS

    )

    if not lines:

        return None

    role_entries = []

    for raw_line in lines:

        original = (raw_line or "").strip()

        if not original:

            continue

        # A bullet is normally a description, not a role heading.

        if re.match(r"^\s*[•●▪◦*+-]\s*", original):

            continue

        line = _clean_entry_line(original)

        if not re.search(r"\b(?:intern|internship|interns)\b", line, re.I):

            continue

        if not _is_likely_title(line, max_words=14):

            continue

        role_entries.append(line)

    count = _distinct_count(role_entries)

    return min(count, 20) if count else None

def extract_programming_language_count(text):

    """

    Count programming languages mentioned in a recognized technical-skills

    section only, not when a language merely appears in project descriptions.

    """

    lines = section_lines(

        text, TECHNICAL_SKILL_HEADINGS,

        COMMON_STOP_HEADINGS - TECHNICAL_SKILL_HEADINGS

    )

    section_text = "\n".join(lines)

    if not section_text.strip():

        return None

    language_patterns = {

        "Python": r"\bpython\b",

        "Java": r"\bjava\b",

        "JavaScript": r"\bjavascript\b",

        "TypeScript": r"\btypescript\b",
        "C++": r"(?<![A-Za-z0-9])C\+\+(?![A-Za-z0-9])",

        "C#": r"(?<![A-Za-z0-9])C#(?![A-Za-z0-9])",

        "C": r"(?<![A-Za-z0-9+#])C(?![A-Za-z0-9+#])",

        "PHP": r"\bphp\b",

        "Ruby": r"\bruby\b",

        "Go": r"\b(?:golang|go language)\b",

        "Kotlin": r"\bkotlin\b",

        "Swift": r"\bswift\b",

        "R": r"(?<![A-Za-z0-9+#])R(?![A-Za-z0-9+#])",

        "SQL": r"\bsql\b",

    }

    found = {

        name for name, pattern in language_patterns.items()

        if re.search(pattern, section_text, re.I)

    }

    return len(found) if found else None

def extract_hackathon_count(text):

    """Count distinct hackathon entries from hackathon-specific sections."""

    lines = section_lines(

        text, HACKATHON_HEADINGS,

        COMMON_STOP_HEADINGS - HACKATHON_HEADINGS

    )

    if not lines:

        return None

    entries = []

    for raw_line in lines:

        line = _clean_entry_line(raw_line)

        if not _is_likely_title(line, max_words=16):

            continue

        # A section heading itself is not included by section_lines().

        entries.append(line)

    count = _distinct_count(entries)

    return min(count, 20) if count else None

def extract_research_paper_count(text):

    """Count actual paper/publication entries, not empty-status statements."""

    lines = section_lines(

        text, RESEARCH_HEADINGS, COMMON_STOP_HEADINGS - RESEARCH_HEADINGS

    )

    if not lines:

        return None

    entries = []

    no_paper_patterns = (

        r"^no\s+(?:research\s+)?papers?\b",

        r"^no\s+publications?\b",

        r"^(?:none|n/?a|not\s+published|no\s+research\s+work)\s*[.!]?$",

        r"^(?:research\s+papers?|publications?)\s*:\s*(?:none|0|n/?a)\s*[.!]?$",

    )

    metadata_pattern = re.compile(

        r"^(?:doi|issn|isbn|url|link|published in|journal|conference|"

        r"authors?|status|under review|available at)\s*:", re.I

    )

    for raw_line in lines:

        line = _clean_entry_line(raw_line)

        if not line:

            continue

        if any(re.search(pattern, line, re.I) for pattern in no_paper_patterns):

            continue

        if metadata_pattern.match(line):

            continue

        if re.search(r"https?://|wwww\.", line, re.I) and len(line.split()) < 4:

            continue

        if not _is_likely_title(line, max_words=30):

            continue

        # Skip explanatory prose; keep concise title-like publication entries.

        if re.search(r"\b(?:this paper|we propose|our work|the study|the research|no papers?)\b", line, re.I) and line.endswith((".", "!", "?")):

            continue

        entries.append(line)

    count = _distinct_count(entries)

    return min(count, 50) if count else None

# RESUME TEXT EXTRACTION

# --------------------------------------------------

def ocr_image(image):

    executable = pytesseract.pytesseract.tesseract_cmd

    if not executable or not os.path.isfile(executable):

        raise RuntimeError(

            "Tesseract OCR was not found. Install Tesseract OCR "

            "or update TESSERACT_PATH in app.py."

        )

    return pytesseract.image_to_string(image)

def extract_pdf_text(file_bytes):

    parts = []

    with pymupdf.open(stream=file_bytes, filetype="pdf") as pdf:

        for page in pdf:

            page_text = page.get_text("text").strip()

            if len(page_text) < 30:

                pixmap = page.get_pixmap(

                    matrix=pymupdf.Matrix(2, 2), alpha=False

                )

                image = Image.open(io.BytesIO(pixmap.tobytes("png")))

                page_text = ocr_image(ImageOps.exif_transpose(image))

            parts.append(page_text)

    return "\n".join(parts)

def extract_docx_text(file_bytes):

    document = Document(io.BytesIO(file_bytes))

    parts = [

        paragraph.text for paragraph in document.paragraphs

        if paragraph.text.strip()

    ]

    for table in document.tables:

        for row in table.rows:

            parts.append(" | ".join(cell.text for cell in row.cells))

    return "\n".join(parts)

def extract_image_text(file_bytes):

    image = Image.open(io.BytesIO(file_bytes))

    image = ImageOps.exif_transpose(image).convert("RGB")

    return ocr_image(image)

def extract_candidate_fields(text):

    """Return resume-derived suggestions; uncertain values remain None."""

    fields = {

        "age": None,

        "education_level": None,

        "university_tier": None,

        "cgpa": None,

        "internships": extract_internship_count(text),

        "projects": extract_project_count(text),

        "programming_languages": extract_programming_language_count(text),

        "certifications": None,

        "experience_years": extract_experience_years(text),

        "hackathons": extract_hackathon_count(text),

        "research_papers": extract_research_paper_count(text),

        "skills_score": calculate_skills_score(text),

        "soft_skills_score": None,

        "resume_length_words": len(text.split()),

        "company_type": None,

    }

    age_match = re.search(r"\bAge\s*[:-]\s*(\d{1,2})\b", text, re.I)

    if age_match:

        age = int(age_match.group(1))

        if 16 <= age <= 100:

            fields["age"] = age

    if re.search(r"\bPh\.?\s?D\b|\bDoctorate\b", text, re.I):

        fields["education_level"] = "PhD"

    elif re.search(

        r"\b(M\.?\s?Tech|M\.?\s?Sc|M\.?\s?E|MBA|Master'?s|Postgraduate)\b",

        text, re.I,

    ):

        fields["education_level"] = "Masters"

    elif re.search(

        r"\b(B\.?\s?Tech|B\.?\s?E|B\.?\s?Sc|BCA|BBA|Bachelor'?s|Undergraduate)\b",

        text, re.I,

    ):

        fields["education_level"] = "Bachelors"

    cgpa_match = re.search(

        r"\b(?:CGPA|C\.?\s?G\.?\s?P\.?\s?A\.?|GPA)\s*[:-]?\s*"

        r"(\d{1,2}(?:\.\d+)?)\s*(?:/\s*10)?\b",

        text, re.I,

    )

    if cgpa_match:

        cgpa = float(cgpa_match.group(1))

        if 0 <= cgpa <= 10:

            fields["cgpa"] = cgpa

    cert_lines = section_lines(

        text,

        CERTIFICATE_HEADINGS,

        COMMON_STOP_HEADINGS - CERTIFICATE_HEADINGS,

    )

    if cert_lines:

        # Support multiple certificates separated by |, semicolon, or bullets.

        entries = []

        for line in cert_lines:

            pieces = re.split(r"\s*[|;]\s*", line.strip())

            for piece in pieces:

                piece = re.sub(r"^\s*(?:[•●▪◦*-]+|\d+[.)])\s*", "", piece).strip()

                if len(piece.split()) >= 2:

                    entries.append(piece)

        unique_entries = []

        seen = set()

        for entry in entries:

            key = re.sub(r"\W+", "", entry.lower())

            if key and key not in seen:

                seen.add(key)

                unique_entries.append(entry)

        if unique_entries:

            fields["certifications"] = len(unique_entries)

    # University tier, soft-skills score, and company type are not guessed.

    return fields

# --------------------------------------------------

# BASIC ROUTES

# --------------------------------------------------

@app.route("/", methods=["GET"])

def home():

    return jsonify({

        "message": "Resume Screening API is running",

        "model": "final_extra_trees",

        "features": len(FEATURES),

    })

@app.route("/model-info", methods=["GET"])

def model_info():

    encoder = preprocessor.named_transformers_["cat"].named_steps["encoder"]

    categorical_options = {

        field: [str(value) for value in categories]

        for field, categories in zip(CATEGORICAL_FIELDS, encoder.categories_)

    }

    return jsonify({

        "model_type": type(model).__name__,

        "feature_count": len(FEATURES),

        "features": FEATURES,

        "numeric_fields": NUMERIC_FIELDS,

        "categorical_fields": CATEGORICAL_FIELDS,

        "categorical_options": categorical_options,

        "classes": [str(value) for value in model.classes_],

    })

# --------------------------------------------------

# SCREENING HISTORY ROUTES

# --------------------------------------------------

@app.route("/history", methods=["GET"])

def get_history():

    connection = get_db_connection()

    try:

        rows = connection.execute("""

            SELECT

                id,

                candidate_data,

                prediction,

                result,

                shortlisted_probability,

                created_at,

                candidate_name,

                resume_filename

            FROM screening_history

            ORDER BY id DESC

        """).fetchall()

        history = []

        for row in rows:

            item = dict(row)

            try:

                item["candidate_data"] = json.loads(

                    item["candidate_data"]

                )

            except (json.JSONDecodeError, TypeError):

                item["candidate_data"] = {}

            item["candidate_name"] = (

                item.get("candidate_name")

                or f"Previous record #{item['id']}"

            )

            item["resume_filename"] = (

                item.get("resume_filename") or ""

            )

            history.append(item)

        return jsonify({

            "history": history,

            "total": len(history)

        }), 200

    except sqlite3.Error:

        app.logger.exception(

            "Could not fetch screening history"

        )

        return jsonify({

            "error": "Unable to fetch screening history."

        }), 500

    finally:

        connection.close()

@app.route("/history/<int:record_id>", methods=["DELETE"])

def delete_history_record(record_id):

    connection = get_db_connection()

    try:

        cursor = connection.execute(

            """

            DELETE FROM screening_history

            WHERE id = ?

            """,

            (record_id,)

        )

        connection.commit()

        if cursor.rowcount == 0:

            return jsonify({

                "error": "History record not found."

            }), 404

        return jsonify({

            "message": "History record deleted successfully.",

            "id": record_id

        }), 200

    except sqlite3.Error:

        connection.rollback()

        app.logger.exception(

            "Could not delete history record"

        )

        return jsonify({

            "error": "Unable to delete history record."

        }), 500

    finally:

        connection.close()

@app.route("/extract-resume", methods=["POST"])

def extract_resume():

    if "resume" not in request.files:

        return jsonify({"error": "No resume file was uploaded."}), 400

    uploaded_file = request.files["resume"]

    if not uploaded_file.filename:

        return jsonify({"error": "Please select a resume file."}), 400

    filename = secure_filename(uploaded_file.filename)

    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if extension not in ALLOWED_EXTENSIONS:

        return jsonify({

            "error": "Unsupported file format.",

            "allowed_formats": sorted(ALLOWED_EXTENSIONS),

        }), 400

    file_bytes = uploaded_file.read()

    if not file_bytes:

        return jsonify({"error": "The uploaded file is empty."}), 400

    try:

        if extension == "pdf":

            text = extract_pdf_text(file_bytes)

        elif extension == "docx":

            text = extract_docx_text(file_bytes)

        else:

            text = extract_image_text(file_bytes)

        text = normalize_text(text)

        if not text:

            return jsonify({

                "error": "No readable text was found in the resume."

            }), 422

        suggestions = extract_candidate_fields(text)

        # Create a history row at upload/extraction time so the uploaded filename
        # is visible in the frontend even before the user runs prediction.
        candidate_name = os.path.splitext(filename)[0].replace("_", " ").strip() or filename
        connection = get_db_connection()
        try:
            cursor = connection.execute(
                """
                INSERT INTO screening_history (
                    candidate_data, prediction, result, shortlisted_probability,
                    created_at, candidate_name, resume_filename
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    json.dumps(suggestions, ensure_ascii=False),
                    "Pending",
                    "Pending",
                    None,
                    datetime.now().isoformat(timespec="seconds"),
                    candidate_name,
                    filename,
                ),
            )
            connection.commit()
            history_id = cursor.lastrowid
        except sqlite3.Error:
            connection.rollback()
            app.logger.exception("Could not save uploaded resume to screening history")
            return jsonify({"error": "Resume was extracted, but its history record could not be saved."}), 500
        finally:
            connection.close()

        return jsonify({
            "message": "Resume text extracted and saved to history successfully.",
            "history_id": history_id,
            "filename": filename,
            "candidate_name": candidate_name,
            "extracted_text": text,
            "suggested_fields": suggestions,
            "candidate_fields": suggestions,
            "review_required": True,
            "note": (
                "Extracted values are suggestions, not verified facts. "
                "Review all fields and manually complete missing details."
            ),
        }), 200

    except Exception:

        app.logger.exception("Resume extraction failed")

        return jsonify({

            "error": (

                "Resume extraction failed. Check backend logs "

                "and the OCR configuration."

            )

        }), 500

# --------------------------------------------------

# SHORTLISTING PREDICTION

# --------------------------------------------------

def is_positive_class(label):

    return str(label).strip().lower() in {

        "1", "1.0", "true", "yes", "shortlisted"

    }

@app.route("/predict", methods=["POST"])

def predict():

    try:

        data = request.get_json(silent=True)

        if not isinstance(data, dict):

            return jsonify({"error": "Expected a JSON request body."}), 400

        missing = [

            field for field in FEATURES

            if field not in data or data[field] is None or data[field] == ""

        ]

        if missing:

            return jsonify({

                "error": "Please provide all 15 model features.",

                "missing_fields": missing,

            }), 400

        candidate = {}

        for field in NUMERIC_FIELDS:

            try:

                value = float(data[field])

            except (ValueError, TypeError):

                return jsonify({

                    "error": f"{field} must be a valid number."

                }), 400

            if not pd.notna(value):

                return jsonify({"error": f"{field} must be a valid number."}), 400

            if value < 0:

                return jsonify({"error": f"{field} cannot be negative."}), 400

            if field in {

                "age", "internships", "projects", "programming_languages",

                "certifications", "hackathons", "research_papers",

                "resume_length_words",

            } and not value.is_integer():

                return jsonify({

                    "error": f"{field} must be a whole number."

                }), 400

            candidate[field] = int(value) if value.is_integer() else value

        if not 16 <= candidate["age"] <= 100:

            return jsonify({"error": "age must be between 16 and 100."}), 400

        if candidate["cgpa"] > 10:

            return jsonify({

                "error": "CGPA cannot exceed 10. Confirm your dataset scale."

            }), 400

        for field in ("skills_score", "soft_skills_score"):

            if not 0 <= candidate[field] <= 100:

                return jsonify({

                    "error": f"{field} must be between 0 and 100."

                }), 400

        encoder = preprocessor.named_transformers_["cat"].named_steps["encoder"]

        category_map = {

            field: {str(value) for value in categories}

            for field, categories in zip(CATEGORICAL_FIELDS, encoder.categories_)

        }

        for field in CATEGORICAL_FIELDS:

            value = data[field]

            if not isinstance(value, str) or not value.strip():

                return jsonify({

                    "error": f"{field} must be a non-empty string."

                }), 400

            value = value.strip()

            if value not in category_map[field]:

                return jsonify({

                    "error": f"Invalid value for {field}.",

                    "allowed_values": sorted(category_map[field]),

                }), 400

            candidate[field] = value

        candidate_df = pd.DataFrame([candidate], columns=FEATURES)

        prediction_value = model.predict(candidate_df)[0]

        classes = list(model.classes_)

        positive_classes = [

            label for label in classes if is_positive_class(label)

        ]

        probability = None

        if hasattr(model, "predict_proba") and positive_classes:

            probabilities = model.predict_proba(candidate_df)[0]

            positive_label = positive_classes[0]

            probability = round(

                float(probabilities[classes.index(positive_label)]) * 100, 2

            )

        result = (
            "Shortlisted"
            if is_positive_class(prediction_value)
            else "Not Shortlisted"
        )

        candidate_name = data.get("candidate_name")
        resume_filename = data.get("resume_filename")
        history_id = data.get("history_id")

        if candidate_name is not None and not isinstance(candidate_name, str):
            return jsonify({"error": "candidate_name must be a string."}), 400
        if resume_filename is not None and not isinstance(resume_filename, str):
            return jsonify({"error": "resume_filename must be a string."}), 400
        if history_id is not None and (isinstance(history_id, bool) or not str(history_id).isdigit()):
            return jsonify({"error": "history_id must be a valid integer."}), 400

        candidate_name = (candidate_name or "").strip() or None
        resume_filename = secure_filename((resume_filename or "").strip()) or None
        connection = get_db_connection()
        try:
            pending_row = None
            if history_id is not None:
                pending_row = connection.execute(
                    "SELECT id, candidate_name, resume_filename FROM screening_history WHERE id = ? AND result = 'Pending'",
                    (int(history_id),),
                ).fetchone()
            elif resume_filename:
                pending_row = connection.execute(
                    "SELECT id, candidate_name, resume_filename FROM screening_history WHERE resume_filename = ? AND result = 'Pending' ORDER BY id DESC LIMIT 1",
                    (resume_filename,),
                ).fetchone()

            if pending_row:
                saved_id = pending_row["id"]
                candidate_name = candidate_name or pending_row["candidate_name"]
                resume_filename = resume_filename or pending_row["resume_filename"]
                connection.execute(
                    """
                    UPDATE screening_history
                    SET candidate_data = ?, prediction = ?, result = ?,
                        shortlisted_probability = ?, created_at = ?,
                        candidate_name = ?, resume_filename = ?
                    WHERE id = ?
                    """,
                    (
                        json.dumps(candidate, ensure_ascii=False),
                        str(prediction_value), result, probability,
                        datetime.now().isoformat(timespec="seconds"),
                        candidate_name, resume_filename, saved_id,
                    ),
                )
            else:
                connection.execute(
                    """
                    INSERT INTO screening_history (
                        candidate_data, prediction, result, shortlisted_probability,
                        created_at, candidate_name, resume_filename
                    ) VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        json.dumps(candidate, ensure_ascii=False),
                        str(prediction_value), result, probability,
                        datetime.now().isoformat(timespec="seconds"),
                        candidate_name, resume_filename,
                    ),
                )
                saved_id = connection.execute("SELECT last_insert_rowid()").fetchone()[0]
            connection.commit()
        except sqlite3.Error:
            connection.rollback()
            app.logger.exception("Could not save screening history")
            return jsonify({
                "error": "Prediction completed, but saving screening history failed."
            }), 500
        finally:
            connection.close()

        return jsonify({
            "message": "Prediction completed and saved successfully.",
            "id": saved_id,
            "candidate_name": candidate_name,
            "resume_filename": resume_filename,
            "prediction": str(prediction_value),
            "result": result,
            "shortlisted_probability": probability,
            "candidate_data": candidate,
        }), 200

    except Exception:

        app.logger.exception("Prediction failed")

        return jsonify({

            "error": "Prediction failed. Check backend logs for details."

        }), 500

# --------------------------------------------------

# ERROR HANDLERS / START SERVER

# --------------------------------------------------

@app.errorhandler(413)

def file_too_large(_error):

    return jsonify({"error": "File exceeds the 10 MB limit."}), 413

if __name__ == "__main__":

    print("Shortlisting model loaded successfully.")

    print("Model type:", type(model).__name__)

    print("Feature count:", len(FEATURES))

    print("Numeric fields:", NUMERIC_FIELDS)

    print("Categorical fields:", CATEGORICAL_FIELDS)

    app.run(host="127.0.0.1", port=5000, debug=False)
